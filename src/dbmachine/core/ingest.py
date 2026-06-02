"""Ingestion: arbitrary CSV / JSON / Excel → entity rows.

Reads a file, infers a column→field mapping heuristically (case/whitespace
-insensitive name matching), and produces rows ready for
:meth:`Runtime.import_rows`. Explicit mappings override the heuristic; LLM-assisted
mapping is deferred (the agent can compute a mapping and pass it in).
"""

from __future__ import annotations

import csv
import json
import re
from pathlib import Path
from typing import Any

from .spec import EntityDef, Spec


def _normalize(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", name.lower())


def read_file(path: str) -> list[dict]:
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"import file not found: {path}")
    suffix = p.suffix.lower()
    if suffix == ".csv":
        return _read_csv(p)
    if suffix in (".json",):
        return _read_json(p)
    if suffix in (".xlsx", ".xlsm"):
        return _read_excel(p)
    raise ValueError(f"unsupported import format {suffix!r} (use .csv, .json, .xlsx)")


def _read_csv(p: Path) -> list[dict]:
    with p.open(newline="") as f:
        return list(csv.DictReader(f))


def _read_json(p: Path) -> list[dict]:
    data = json.loads(p.read_text())
    if isinstance(data, dict):
        # allow {"rows": [...]} or a single object
        data = data.get("rows", [data])
    if not isinstance(data, list):
        raise ValueError("JSON import must be an array of objects (or {rows: [...]})")
    return data


def _read_excel(p: Path) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(p, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = [str(h) if h is not None else "" for h in next(rows_iter)]
    except StopIteration:
        return []
    out = []
    for row in rows_iter:
        if all(c is None for c in row):
            continue
        out.append({header[i]: row[i] for i in range(min(len(header), len(row)))})
    return out


def entity_targets(spec: Spec, entity: str) -> list[str]:
    """Importable target columns for an entity (fields + belongs_to *_id)."""
    e: EntityDef = spec.entities[entity]
    targets = list(e.fields.keys())
    targets += [f"{r}_id" for r, rel in e.relationships.items() if rel.kind == "belongs_to"]
    return targets


def infer_mapping(targets: list[str], source_columns: list[str]) -> dict[str, str]:
    """Map source column → target field by normalized-name match."""
    by_norm = {_normalize(t): t for t in targets}
    mapping: dict[str, str] = {}
    for src in source_columns:
        norm = _normalize(src)
        if norm in by_norm:
            mapping[src] = by_norm[norm]
    return mapping


def prepare(
    spec: Spec,
    entity: str,
    path: str,
    *,
    mapping: dict[str, str] | None = None,
) -> dict[str, Any]:
    """Read + map a file into rows for import.

    Returns a dict with ``rows``, the resolved ``mapping``, and any
    ``unmapped_source`` / ``unmapped_target`` columns for the agent to inspect.
    """
    if entity not in spec.entities:
        raise ValueError(f"unknown entity {entity!r}")
    raw = read_file(path)
    source_columns = list(raw[0].keys()) if raw else []
    targets = entity_targets(spec, entity)

    resolved = dict(infer_mapping(targets, source_columns))
    if mapping:
        resolved.update(mapping)  # explicit overrides heuristic

    rows = []
    for r in raw:
        mapped = {}
        for src, dest in resolved.items():
            if src in r and r[src] not in (None, ""):
                mapped[dest] = r[src]
        if mapped:
            rows.append(mapped)

    return {
        "entity": entity,
        "rows": rows,
        "mapping": resolved,
        "unmapped_source": [c for c in source_columns if c not in resolved],
        "unmapped_target": [t for t in targets if t not in resolved.values()],
        "row_count": len(rows),
    }
